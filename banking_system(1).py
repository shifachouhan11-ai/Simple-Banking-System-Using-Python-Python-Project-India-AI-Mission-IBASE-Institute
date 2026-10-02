import openpyxl
import os


file_name = "banking_system.xlsx"


# Create Excel File
def create_file():

    if os.path.exists(file_name):
        return

    workbook = openpyxl.Workbook()
    sheet = workbook.active

    sheet.append(["Account Number", "Name", "Mobile Number", "PIN", "Balance"])

    workbook.save(file_name)


# Find Account
def find_account(account_number):

    workbook = openpyxl.load_workbook(file_name)
    sheet = workbook.active

    for row in range(2, sheet.max_row + 1):

        if sheet.cell(row, 1).value == account_number:
            workbook.close()
            return row

    workbook.close()
    return 0


# Verify PIN
def verify_pin(row, pin):

    workbook = openpyxl.load_workbook(file_name)
    sheet = workbook.active

    saved_pin = sheet.cell(row, 4).value

    workbook.close()

    if saved_pin == pin:
        return True
    else:
        return False


# Create Account
def create_account():

    print("\n----- CREATE ACCOUNT -----")

    name = input("Enter Account Holder Name: ")

    mobile = input("Enter Mobile Number: ")

    pin = input("Enter 4 Digit PIN: ")

    balance = float(input("Enter Initial Deposit: "))

    workbook = openpyxl.load_workbook(file_name)
    sheet = workbook.active

    if sheet.max_row == 1:
        account_number = 1001
    else:
        last_account = sheet.cell(sheet.max_row, 1).value
        account_number = last_account + 1

    sheet.append([account_number, name, mobile, pin, balance])

    workbook.save(file_name)
    workbook.close()

    print("\nAccount Created Successfully!")
    print("Your Account Number is:", account_number)


# Deposit Money
def deposit_money():

    print("\n----- DEPOSIT MONEY -----")

    account_number = int(input("Enter Account Number: "))

    row = find_account(account_number)

    if row == 0:
        print("Account not found!")
        return

    pin = input("Enter PIN: ")

    if verify_pin(row, pin) == False:
        print("Wrong PIN!")
        return

    amount = float(input("Enter Deposit Amount: "))

    if amount <= 0:
        print("Invalid Amount!")
        return

    workbook = openpyxl.load_workbook(file_name)
    sheet = workbook.active

    balance = sheet.cell(row, 5).value

    balance = balance + amount

    sheet.cell(row, 5).value = balance

    workbook.save(file_name)
    workbook.close()

    print("Money Deposited Successfully!")
    print("Updated Balance:", balance)


# Withdraw Money
def withdraw_money():

    print("\n----- WITHDRAW MONEY -----")

    account_number = int(input("Enter Account Number: "))

    row = find_account(account_number)

    if row == 0:
        print("Account not found!")
        return

    pin = input("Enter PIN: ")

    if verify_pin(row, pin) == False:
        print("Wrong PIN!")
        return

    amount = float(input("Enter Withdrawal Amount: "))

    if amount <= 0:
        print("Invalid Amount!")
        return

    workbook = openpyxl.load_workbook(file_name)
    sheet = workbook.active

    balance = sheet.cell(row, 5).value

    if amount > balance:
        print("Insufficient Balance!")

        workbook.close()
        return

    balance = balance - amount

    sheet.cell(row, 5).value = balance

    workbook.save(file_name)
    workbook.close()

    print("Money Withdrawn Successfully!")
    print("Updated Balance:", balance)


# Check Balance
def check_balance():

    print("\n----- CHECK BALANCE -----")

    account_number = int(input("Enter Account Number: "))

    row = find_account(account_number)

    if row == 0:
        print("Account not found!")
        return

    pin = input("Enter PIN: ")

    if verify_pin(row, pin) == False:
        print("Wrong PIN!")
        return

    workbook = openpyxl.load_workbook(file_name)
    sheet = workbook.active

    name = sheet.cell(row, 2).value
    balance = sheet.cell(row, 5).value

    workbook.close()

    print("\nAccount Number:", account_number)
    print("Account Holder Name:", name)
    print("Current Balance:", balance)


# Main Function
def main():

    create_file()

    while True:

        print("\n==========================")
        print("   SIMPLE BANKING SYSTEM")
        print("==========================")
        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Check Balance")
        print("5. Exit")
        print("==========================")

        choice = input("Enter your choice: ")

        if choice == "1":
            create_account()

        elif choice == "2":
            deposit_money()

        elif choice == "3":
            withdraw_money()

        elif choice == "4":
            check_balance()

        elif choice == "5":
            print("Thank you for using Banking System!")
            break

        else:
            print("Invalid Choice!")


# Start Program
main()
